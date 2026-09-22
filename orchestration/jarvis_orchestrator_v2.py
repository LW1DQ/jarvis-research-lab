#!/usr/bin/env python3
"""
JARVIS LangGraph Orchestrator v2
Adapted from Phoenix1454/Multi-Agent-Research-Assistant-Langgraph
for local Ollama + MCP servers + PaperQA2 + NS-3
with LangFuse observability (when available)
"""

import os
import json
import asyncio
import logging
from datetime import datetime
from typing import TypedDict, Annotated, List, Dict, Any, Optional
from langgraph.graph import StateGraph, END
from langgraph.checkpoint.memory import MemorySaver
import operator

# ============================================================
# SETUP: Environment + Native Ollama Client
# ============================================================

# CRITICAL: Must be before any imports that use OPENAI_API_KEY
for var in ['AGENT', 'OPENAI_API_KEY', 'OPENAI_API_BASE', 'OPENROUTER_API_KEY', 'OPENAI_ADMIN_KEY', 'OPENAI_ORGANIZATION']:
    os.environ.pop(var, None)

os.environ['OPENAI_API_KEY'] = 'EMPTY'
os.environ['OPENAI_API_BASE'] = 'http://127.0.0.1:11434/v1'
os.environ['OLLAMA_API_BASE'] = 'http://127.0.0.1:11434'

# LangFuse Integration (Optional - will work when LangFuse is available)
LANGFUSE_AVAILABLE = False
LANGFUSE_HANDLER = None
try:
    from langfuse import Langfuse
    from langfuse.callback import CallbackHandler
    
    LANGFUSE_PUBLIC_KEY = os.getenv("LANGFUSE_PUBLIC_KEY")
    LANGFUSE_SECRET_KEY = os.getenv("LANGFUSE_SECRET_KEY")
    LANGFUSE_HOST = os.getenv("LANGFUSE_HOST", "http://localhost:3000")
    
    if LANGFUSE_PUBLIC_KEY and LANGFUSE_SECRET_KEY:
        langfuse = Langfuse(
            public_key=LANGFUSE_PUBLIC_KEY,
            secret_key=LANGFUSE_SECRET_KEY,
            host=LANGFUSE_HOST
        )
        LANGFUSE_HANDLER = CallbackHandler(langfuse=langfuse)
        LANGFUSE_AVAILABLE = True
        print("✅ LangFuse integration enabled")
    else:
        print("ℹ️  LangFuse not configured (set LANGFUSE_PUBLIC_KEY and LANGFUSE_SECRET_KEY)")
except ImportError:
    print("ℹ️  LangFuse not installed (pip install langfuse)")

# Native Ollama client
import ollama
_OLLAMA_CLIENT = ollama.Client(host='http://127.0.0.1:11434')

# ============================================================
# LLM CONFIGURATIONS - qwen2.5-coder:7b for fast responses, qwen3:8b for reasoning
# ============================================================

LLM_CONFIG = {
    "model": "qwen2.5-coder:7b",
    "options": {"temperature": 0.0, "num_predict": 4096, "think": False},
}

REASONING_LLM_CONFIG = {
    "model": "qwen3:8b",
    "options": {"temperature": 0.0, "num_predict": 4096, "think": False},
}

CODING_LLM_CONFIG = {
    "model": "qwen2.5-coder:7b",
    "options": {"temperature": 0.0, "num_predict": 4096, "think": False},
}

# Nemotron 3 Ultra Free (NVIDIA) - Cloud premium reasoning/coding
NEMOTRON_ULTRA_FREE_CONFIG = {
    "model": "nvidia/nemotron-3-ultra-550b-a55b",
    "api_base": "https://integrate.api.nvidia.com/v1",
    "options": {"temperature": 0.0, "num_predict": 4096},
}

# ============================================================
# FILE-BASED OBSERVABILITY (JSONL Logging) - Always active
# ============================================================

OBS_DIR = "/home/diego/research-jarvis/logs/observability"
os.makedirs(OBS_DIR, exist_ok=True)

OBS_FILE = os.path.join(OBS_DIR, f"jarvis_{datetime.now().strftime('%Y%m%d')}.jsonl")

def log_obs(event: str, data: Dict[str, Any], level: str = "INFO"):
    """Log structured observability event to JSONL file."""
    entry = {
        "timestamp": datetime.now().isoformat(),
        "level": level,
        "event": event,
        "data": data
    }
    with open(OBS_FILE, "a") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")

# Configure standard logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger("jarvis")

# ============================================================
# LLM CALL HELPER (Native Ollama)
# ============================================================

async def call_ollama(config: Dict[str, Any], prompt: str, system_prompt: str = None) -> str:
    """Call Ollama native API with observability."""
    start_time = datetime.now()
    messages = []
    if system_prompt:
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": prompt})
    
    try:
        resp = _OLLAMA_CLIENT.chat(
            model=config["model"],
            messages=messages,
            options=config["options"]
        )
        content = resp["message"]["content"]
        log_obs("ollama_call", {
            "model": config["model"],
            "duration_ms": (datetime.now() - start_time).total_seconds() * 1000,
            "prompt_len": len(prompt),
            "response_len": len(content)
        }, "INFO")
        return content
    except Exception as e:
        log_obs("ollama_call", {"error": str(e), "duration_ms": (datetime.now() - start_time).total_seconds() * 1000}, "ERROR")
        return f"Error: {e}"

async def call_nemotron(config: Dict[str, Any], prompt: str, system_prompt: str = None) -> str:
    """Call Nemotron 3 Ultra Free API (NVIDIA) with observability."""
    start_time = datetime.now()
    messages = []
    if system_prompt:
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": prompt})
    
    api_key = os.getenv("NEMOTRON_API_KEY")
    if not api_key:
        log_obs("nemotron_call", {"error": "NEMOTRON_API_KEY not set"}, "ERROR")
        return "Error: NEMOTRON_API_KEY not set in environment"
    
    try:
        import httpx
        async with httpx.AsyncClient(timeout=60.0) as client:
            headers = {
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json"
            }
            payload = {
                "model": config["model"],
                "messages": messages,
                "temperature": config["options"].get("temperature", 0.0),
                "max_tokens": config["options"].get("num_predict", 4096)
            }
            resp = await client.post(
                f"{config['api_base']}/chat/completions",
                headers=headers,
                json=payload
            )
            resp.raise_for_status()
            data = resp.json()
            content = data["choices"][0]["message"]["content"]
            log_obs("nemotron_call", {
                "model": config["model"],
                "duration_ms": (datetime.now() - start_time).total_seconds() * 1000,
                "prompt_len": len(prompt),
                "response_len": len(content)
            }, "INFO")
            return content
    except Exception as e:
        log_obs("nemotron_call", {"error": str(e), "duration_ms": (datetime.now() - start_time).total_seconds() * 1000}, "ERROR")
        return f"Error: {e}"

# ============================================================
# LLM CALL HELPER WITH LANGFUSE (when available)
# ============================================================

async def call_llm(prompt: str, event_name: str, metadata: Dict = None, config: Dict = None) -> str:
    """Call LLM with observability logging + LangFuse tracing."""
    if config is None:
        config = LLM_CONFIG
    
    start_time = datetime.now()
    
    # Prepare LangFuse metadata
    langfuse_metadata = {
        "event_name": event_name,
        "metadata": metadata or {},
        "tags": ["jarvis", "orchestrator"]
    }
    
    callbacks = [LANGFUSE_HANDLER] if LANGFUSE_HANDLER else None
    
    try:
        # Use native Ollama client
        resp = _OLLAMA_CLIENT.chat(
            model=config["model"],
            messages=[{"role": "user", "content": prompt}],
            options=config["options"]
        )
        content = resp["message"]["content"]
        
        # Log to JSONL
        log_obs("llm_call", {
            "event": event_name,
            "model": config["model"],
            "duration_ms": (datetime.now() - start_time).total_seconds() * 1000,
            "prompt_len": len(prompt),
            "response_len": len(content),
            "metadata": metadata or {}
        }, "INFO")
        
        # Log to LangFuse if available
        if LANGFUSE_HANDLER:
            try:
                trace = langfuse.trace(
                    name=f"jarvis_{event_name}",
                    input=prompt,
                    output=content,
                    metadata={**(metadata or {}), "model": config["model"]}
                )
            except:
                pass
        
        return content
    except Exception as e:
        log_obs("llm_call", {"event": event_name, "error": str(e)}, "ERROR")
        return f"Error: {e}"

# ============================================================
# MCP CLIENTS
# ============================================================

import httpx

class MCPClient:
    """Simple HTTP client for FastMCP servers (JSON-RPC 2.0 over SSE)."""
    
    def __init__(self, base_url: str, name: str):
        self.base_url = base_url
        self.name = name
        self.client = httpx.AsyncClient(timeout=60.0)
        self.session_id: Optional[str] = None
    
    async def initialize(self):
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.post(
                f"{self.base_url}/mcp",
                headers={'Content-Type': 'application/json', 'Accept': 'application/json, text/event-stream'},
                json={'jsonrpc':'2.0','id':1,'method':'initialize','params':{'protocolVersion':'2024-11-05','capabilities':{},'clientInfo':{'name':'jarvis','version':'1.0'}}}
            )
            self.session_id = resp.headers.get('mcp-session-id')
            return self.session_id
    
    async def call_tool(self, tool_name: str, arguments: Dict[str, Any]) -> Any:
        if not self.session_id:
            await self.initialize()
        
        async with httpx.AsyncClient(timeout=60.0) as client:
            resp = await client.post(
                f"{self.base_url}/mcp",
                headers={
                    'Content-Type': 'application/json', 
                    'Accept': 'application/json, text/event-stream', 
                    'mcp-session-id': self.session_id
                },
                json={'jsonrpc':'2.0','id':2,'method':'tools/call','params':{'name':tool_name,'arguments':arguments}}
            )
            
            for line in resp.text.split('\n'):
                if line.startswith('data: '):
                    try:
                        data = json.loads(line[6:])
                        if 'result' in data:
                            result = data['result']
                            if isinstance(result, dict) and 'content' in result:
                                return result['content'][0]['text'] if result['content'] else result
                            return result
                    except Exception as e:
                        pass
            return "No result"
    
    async def close(self):
        await self.client.aclose()

MCP_SERVERS = {
    "research": {"url": "http://127.0.0.1:8001", "name": "research-mcp"},
    "ns3": {"url": "http://127.0.0.1:8002", "name": "ns3-mcp"},
    "python": {"url": "http://127.0.0.1:8003", "name": "python-mcp"},
    "personal": {"url": "http://127.0.0.1:8004", "name": "personal-mcp"},
}

_mcp_clients: Dict[str, MCPClient] = {}

async def get_mcp_client(name: str) -> MCPClient:
    global _mcp_clients
    if name not in _mcp_clients:
        cfg = MCP_SERVERS[name]
        _mcp_clients[name] = MCPClient(cfg["url"], cfg["name"])
        await _mcp_clients[name].initialize()
    return _mcp_clients[name]

async def close_mcp_clients():
    for client in _mcp_clients.values():
        await client.close()
    _mcp_clients.clear()

# ============================================================
# STATE DEFINITION
# ============================================================

class JarvisState(TypedDict):
    """State for the JARVIS research workflow."""
    main_task: str
    research_findings: Annotated[List[str], operator.add]
    ns3_results: Annotated[List[str], operator.add]
    draft: str
    critique_notes: str
    revision_number: int
    next_step: str
    current_sub_task: str
    metadata: Dict[str, Any]

# ============================================================
# AGENT NODES (Phoenix1454 Pattern)
# ============================================================

# LLM configurations
LLM_CONFIG = {"model": "qwen3:8b", "options": {"temperature": 0.0, "num_predict": 4096}, "think": False}
CODING_LLM_CONFIG = {"model": "qwen2.5-coder:7b", "options": {"temperature": 0.0, "num_predict": 4096}}

# ----------------- #
# SUPERVISOR NODE   #
# ----------------- #

async def supervisor_node(state: JarvisState) -> Dict[str, Any]:
    research = state.get("research_findings", [])
    has_research = len(research) > 0
    has_draft = bool(state.get("draft", "").strip())
    critique = state.get("critique_notes", "")
    revision = state.get("revision_number", 0)
    
    # 1. If critique says APPROVED, we're done
    if "APPROVED" in critique.upper() and state.get("draft", "").strip():
        return {"next_step": "END", "current_sub_task": "Report approved and complete"}
    
    # 2. If no research yet, start with research
    if not state.get("research_findings"):
        return {"next_step": "researcher", "current_sub_task": f"Research the topic: {state.get('main_task', '')}"}
    
    # 3. If we have research but no draft, create first draft
    if has_research and not has_draft:
        return {"next_step": "writer", "current_sub_task": "Write the first draft based on research findings"}
    
    # 4. If we have a draft but no critique yet, send to critiquer
    if has_draft and not state.get("critique_notes"):
        return {"next_step": "critiquer", "current_sub_task": "Critique the draft"}
    
    # 5. If we have critique with feedback (not approved), revise
    if critique and "APPROVED" not in critique.upper() and revision < 3:
        return {"next_step": "writer", "current_sub_task": "Revise the draft based on critique feedback"}
    
    # Max revisions reached
    if revision >= 3:
        return {"next_step": "END", "current_sub_task": "Maximum revisions reached"}
    
    return {"next_step": "writer", "current_sub_task": "Continue with draft creation"}

# ----------------- #
# RESEARCHER NODE   #
# ----------------- #

async def researcher_node(state: JarvisState) -> Dict[str, Any]:
    sub_task = state.get("current_sub_task", state.get("main_task", ""))
    findings = []
    
    # 1. Search arXiv via research-mcp
    try:
        client = await get_mcp_client("research")
        result = await client.call_tool("search_arxiv", {
            "query": state.get("main_task", "") + " " + state.get("current_sub_task", ""),
            "max_results": 5
        })
        findings.append(f"arXiv:\n{result}")
    except Exception as e:
        pass
    
    # 2. Search OpenAlex
    try:
        client = await get_mcp_client("research")
        result = await client.call_tool("search_openalex", {
            "query": state.get("main_task", ""),
            "max_results": 5
        })
        findings.append(f"OpenAlex:\n{result}")
    except Exception as e:
        pass
    
    # 3. Search Crossref
    try:
        client = await get_mcp_client("research")
        result = await client.call_tool("search_crossref", {
            "query": state.get("main_task", ""),
            "max_results": 5
        })
        findings.append(f"Crossref:\n{result}")
    except Exception as e:
        pass
    
    # 4. Synthesize findings with LLM
    research_text = "\n\n".join(findings) if findings else "No external research found."
    
    synthesis_prompt = f"""Based on the following research about "{state.get('main_task', '')}", provide a concise synthesis of key findings (5-7 bullet points):

{research_text}

Format as clear bullet points with the most important information."""
    
    summary = await call_llm(synthesis_prompt, "researcher_synthesis", {"sub_task": state.get("current_sub_task", "")})
    findings.append(f"Synthesis:\n{summary}")
    
    # Save to personal memory
    try:
        client = await get_mcp_client("personal")
        await client.call_tool("add_note", {
            "title": f"Research: {state.get('main_task', '')[:50]}",
            "content": "\n\n".join(findings),
            "tags": "research,auto"
        })
    except:
        pass
    
    return {"research_findings": findings}

# ----------------- #
# WRITER NODE       #
# ----------------- #

async def writer_node(state: JarvisState) -> Dict[str, Any]:
    research = state.get("research_findings", [])
    research_text = "\n\n".join(research) if research else "No research available."
    ns3_results = state.get("ns3_results", [])
    ns3_text = "\n\n".join(ns3_results) if ns3_results else "No simulation results."
    
    prompt = f"""You are a professional technical writer producing a research report for a PhD thesis in network simulation (NS-3 + Deep Learning).

Report Brief: {state.get('main_task', '')}

Research Material:
{research_text}

NS-3 Simulation Results:
{ns3_text}

Previous Draft (if any):
{state.get('draft', 'No previous draft.')}

Editor Feedback (if any):
{state.get('critique_notes', 'No critique yet.')}

Revision Number: {state.get('revision_number', 0)}

Writing guidelines:
- Structure: Executive Summary → Background → Methodology → Results → Discussion → Conclusion
- Tone: Academic, precise, suitable for PhD thesis
- Length: 1500-2500 words
- Ground every claim in the provided research — no hallucination
- If revision, address every point raised in critique directly
- Use subheadings, short paragraphs, bullet lists where they improve readability
- End with 3-5 concrete actionable takeaways

Write the complete report now:"""
    
    content = await call_llm(prompt, "writer_draft", {"revision": state.get("revision_number", 0)})
    
    # Generate DOCX via python-mcp
    try:
        client = await get_mcp_client("python")
        code = f'''
from docx import Document
doc = Document()
doc.add_heading("JARVIS Research Report", 0)
doc.add_paragraph("Generated by Writer Agent at {datetime.now().isoformat()}")
for para in """{content}""".split("\\n\\n"):
    if para.strip():
        if para.startswith("# "):
            doc.add_heading(para[2:], 1)
        elif para.startswith("## "):
            doc.add_heading(para[3:], 2)
        else:
            doc.add_paragraph(para)
doc.save("/home/diego/research-jarvis/projects/JARVIS_Report_{{datetime.now().strftime('%Y%m%d_%H%M%S')}}.docx")
print("DOCX saved")
'''
        await client.call_tool("run_python", {"code": code})
    except:
        pass
    
    return {"draft": content, "revision_number": state.get("revision_number", 0) + 1}

# ----------------- #
# NS-3 SIMULATION NODE #
# ----------------- #

async def ns3_node(state: JarvisState) -> Dict[str, Any]:
    findings = []
    try:
        client = await get_mcp_client("ns3")
        result = await client.call_tool("list_ns3_modules", {})
        if "result" in result:
            findings.append(f"NS-3 modules: {result['result']}")
        result = await client.call_tool("ns3_ai_status", {})
        if result:
            findings.append(f"ns3-ai status: {result}")
    except:
        findings = ["NS-3 simulation error"]
    return {"ns3_results": findings}

# ----------------- #
# CRITIQUE NODE     #
# ----------------- #

async def critiquer_node(state: JarvisState) -> Dict[str, Any]:
    draft = state.get("draft", "")
    revision_num = state.get("revision_number", 0)
    
    if len(draft.strip()) < 200:
        return {"critique_notes": "APPROVED - Draft is minimal but acceptable", "next_step": "END"}
    if state.get("revision_number", 0) >= 3:
        return {"critique_notes": "APPROVED - Maximum revisions reached", "next_step": "END"}
    
    prompt = f"""You are a rigorous editorial reviewer for a PhD thesis in network simulation (NS-3 + Deep Learning).

Report Topic: {state.get('main_task', '')}

Draft Under Review:
{state.get('draft', '')}

Evaluate the draft across these five dimensions:

1. Coverage — Does it address the core topic without major gaps?
2. Evidence — Are claims backed by research data and NS-3 results, not just assertions?
3. Structure — Executive Summary → Background → Methodology → Results → Discussion → Conclusion in logical order?
4. Clarity — Are technical concepts explained without unnecessary jargon?
5. Actionability — Does the reader leave with something useful they can act on?

Decision:
- If the draft scores well on all five: respond with "APPROVED - [one sentence on its strongest quality]"
- If improvements are needed: list specific, numbered, actionable fixes for the writer

Your review:"""
    
    content = await call_llm(prompt, "critique_review", {"revision": state.get("revision_number", 0)})
    is_approved = "APPROVED" in content.upper()
    
    if is_approved:
        return {"critique_notes": content, "next_step": "END"}
    else:
        return {"critique_notes": content, "next_step": "writer"}

# ============================================================
# BUILD THE GRAPH
# ============================================================

def build_jarvis_graph():
    workflow = StateGraph(JarvisState)
    workflow.add_node("supervisor", supervisor_node)
    workflow.add_node("researcher", researcher_node)
    workflow.add_node("writer", writer_node)
    workflow.add_node("ns3_simulation", ns3_node)
    workflow.add_node("critiquer", critiquer_node)
    workflow.set_entry_point("supervisor")
    workflow.add_edge("researcher", "supervisor")
    workflow.add_edge("writer", "critiquer")
    workflow.add_edge("ns3_simulation", "supervisor")
    
    workflow.add_conditional_edges(
        "critiquer",
        lambda state: state.get("next_step", "writer"),
        {"writer": "writer", "END": END}
    )
    
    workflow.add_conditional_edges(
        "supervisor",
        lambda state: state.get("next_step", "researcher"),
        {"researcher": "researcher", "writer": "writer", "critiquer": "critiquer", "ns3_simulation": "ns3_simulation", "END": END}
    )
    
    app = workflow.compile(checkpointer=MemorySaver())
    return app

# ============================================================
# MAIN ENTRY POINT
# ============================================================

async def run_jarvis_research(main_task: str) -> Dict[str, Any]:
    app = build_jarvis_graph()
    initial_state = {
        "main_task": main_task,
        "research_findings": [], "ns3_results": [], "draft": "",
        "critique_notes": "", "revision_number": 0,
        "next_step": "supervisor", "current_sub_task": "",
        "metadata": {"start_time": datetime.now().isoformat()}
    }
    
    timestamp_fmt = '%Y%m%d_%H%M%S'
    config = {"configurable": {"thread_id": f"jarvis_{datetime.now().strftime(timestamp_fmt)}"}}
    result = await app.ainvoke(initial_state, config=config)
    return result

# ============================================================
# CLI ENTRY POINT
# ============================================================

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        task = " ".join(sys.argv[1:])
    else:
        task = "Investigate MARL for WiFi scheduling in dense networks using NS-3 simulations"
    
    print(f"🚀 JARVIS Research v2: {task}")
    try:
        result = asyncio.run(run_jarvis_research(task))
        print(f"\n📊 FINAL: Revisions={result.get('revision_number',0)}, Sources={len(result.get('research_findings',[]))}, Draft={len(result.get('draft',''))} chars")
    except KeyboardInterrupt:
        print("\n⚠️ Interrupted")
    except Exception as e:
        print(f"\n❌ Error: {e}")
        import traceback; traceback.print_exc()
    finally:
        asyncio.run(close_mcp_clients())
